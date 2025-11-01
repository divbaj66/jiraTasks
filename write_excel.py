# write_excel.py

from openpyxl import load_workbook

def write_to_excel(excel_path, data_rows, sheet_name="DOD Staging"):
    workbook = load_workbook(excel_path)
    sheet = workbook[sheet_name]
    
    # ✅ Step 1: Clear all rows except header
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row)

    for row in data_rows:
        sheet.append(row)

    workbook.save(excel_path)

    print(f"✅ Data written to {excel_path}")
